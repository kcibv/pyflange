'''Extract the SGRE reference values from the validation workbooks into JSON files.

This is a one-off conversion tool that documents where the reference data used by
`validate.py` comes from. It needs `openpyxl` (not a pyflange dependency)
and is only required when the reference workbooks change.

For each validation workbook it writes a case file `references/<case-id>.json`
containing:

- the flange segment input parameters (taken from the scripts that originally
  generated the PyFlange results: `gen-cases.py` and `Case-D7500_T.py`),
- the fatigue analysis parameters,
- the SGRE reference values for each gap angle, read from the `SGRE` sheet
  (L-flange workbooks) or from the `AMD1` sheet (T-flange workbooks).

The Markov matrices are written to `references/markov/` and shared by the cases
that use the same one.

When a case file already exists, its `tolerances` object is preserved.

Usage:
    py validations/extract_references.py
'''

import os
import json
import hashlib

import openpyxl


HERE = os.path.dirname(os.path.abspath(__file__))
L_DIR = os.path.join(HERE, "flangesegments.PolynomialLFlangeSegment")
T_DIR = os.path.join(HERE, "flangesegments.PolynomialTFlangeSegment")
OUT_DIR = os.path.join(HERE, "references")

mm = 0.001
kN = 1000
MPa = 1e6

GAP_ANGLES = [30, 60, 90, 120]   # deg



# ------------------------------------------------------------------------------
#  Compared quantities
#
#  Each entry maps a quantity key (as understood by validate.py) to the
#  row of the SGRE sheet holding its reference value, plus a symbol and a unit
#  for reporting. The selection and the rows reproduce the 'Comparison' sheet
#  of the L-flange workbooks and the 'AMD1' sheet of the T-flange workbooks.
# ------------------------------------------------------------------------------

COMMON_QUANTITIES = {
    "bolt.fy":              (29,  "fy,b",               "MPa"),
    "bolt.As":              (30,  "AS",                 "mm²"),
    "bolt.Fv":              (31,  "Fp,C'",              "kN"),
    "bolt.dS":              (36,  "dS",                 "m/N"),
    "flange.dP":            (37,  "dP",                 "m/N"),
    "p":                    (38,  "p",                  "-"),
    "gap.angle":            (48,  "aGap",               "deg"),
    "gap.L":                (49,  "LGap",               "m"),
    "D":                    (50,  "D",                  "m"),
    "gap.k_mean":           (52,  "kmean",              "mm"),
    "gap.COV":              (53,  "COVk",               "-"),
    "gap.height":           (60,  "kdesign",            "mm"),
    "Z1":                   (67,  "Z1",                 "kN"),
    "Fs1":                  (68,  "FS,1",               "kN"),
    "Z2":                   (69,  "Z2",                 "kN"),
    "Fs2":                  (70,  "FS,2",               "kN"),
    "Z3":                   (71,  "Z3",                 "kN"),
    "Fs3":                  (72,  "FS,3",               "kN"),
    "Fs_poly.a2":           (76,  "a2",                 "1/kN"),
    "Fs_poly.a1":           (77,  "a1",                 "-"),
    "Fs_poly.a0":           (78,  "a0",                 "kN"),
    "DZ_gap":               (79,  "DZgap",              "kN"),
    "I_tg":                 (81,  "Itg",                "mm4"),
    "A_cf":                 (82,  "Acf",                "mm²"),
    "I_cf":                 (83,  "Icf",                "mm4"),
    "k_gap_tot":            (84,  "kgap,tot",           "N/mm/mm"),
    "f_tot":                (85,  "ftot",               "-"),
    "k_shell":              (88,  "kshell,ini",         "N/mm/mm"),
    "k_fac":                (89,  "kfac",               "-"),
    "k_fl":                 (90,  "kfl",                "N/mm/mm"),
    "DZ_gap_inclination":   (93,  "DZgap,inclination",  "kN"),
    "DZ_gap_tot":           (94,  "DZgap,tot",          "kN"),
    "Z0":                   (96,  "Z0",                 "kN"),
    "ak":                   (98,  "ak",                 "-"),
    "k_seg":                (100, "kseg",               "N/mm/mm"),
    "Z_dw":                 (105, "DZdw",               "kN"),
    "Z_close":              (107, "Zclose",             "kN"),
    "c_ini_true":           (108, "cini,true",          "-"),
    "Fs_min":               (109, "FS,min",             "kN"),
    "a_star":               (113, "a*",                 "mm"),
    "bolt.bS":              (114, "bS",                 "1/Nmm"),
    "damage":               (136, "Damage",             "-"),
}

L_QUANTITIES = {
    **COMMON_QUANTITIES,
    "u":                    (102, "u",                  "mm"),
    "c_ini_mod":            (110, "cini,mod",           "-"),
    "c_ini_M":              (118, "cini,M",             "-"),
    "Ms_min":               (120, "MS,min",             "Nm"),
}

T_QUANTITIES = {
    **COMMON_QUANTITIES,
    "DZ_gap_c":             (91,  "DZgap,c",            "kN"),
    "Fv_c":                 (92,  "FV,c",               "kN"),
    "u":                    (103, "u",                  "mm"),
    "c_ini_mod":            (111, "cini,mod",           "-"),
    "c_ini_M":              (119, "cini,M",             "-"),
    "Ms_min":               (121, "MS,min",             "Nm"),
    "Ms_T.at_rest":         (125, "Ms,T",               "Nm"),
    "Ms_T.at_rest+100N":    (128, "Ms,T",               "Nm"),
    "M1":                   (129, "M1",                 "Nm"),
    "M2":                   (130, "M2",                 "Nm"),
    "M3":                   (131, "M3",                 "Nm"),
    "Ms_poly.a2":           (132, "a2,M",               "1/kN²·Nm"),
    "Ms_poly.a1":           (133, "a1,M",               "Nm/kN"),
    "Ms_poly.a0":           (134, "a0,M",               "Nm"),
}



# ------------------------------------------------------------------------------
#  Case input parameters (SI units: m, N, Pa, rad)
# ------------------------------------------------------------------------------

M80 = dict(
    nominal_diameter = 80*mm,
    thread_pitch = 6*mm,
    shank_diameter_ratio = 76.1/80,
    shank_length = 270*mm,
    yield_stress = 900*MPa,
    ultimate_tensile_stress = 1000*MPa,
    stud = True)

M80_hex_nut = dict(
    nominal_diameter = 80*mm,
    thickness = 64*mm,
    inscribed_diameter = 115*mm,
    circumscribed_diameter = 127.5*mm,
    bearing_diameter = 140*mm)

M48_L = dict(
    nominal_diameter = 48*mm,
    thread_pitch = 5*mm,
    shank_length = 238*mm,
    yield_stress = 900*MPa,
    ultimate_tensile_stress = 1040*MPa,
    stud = False)

M48_L_hex_nut = dict(
    nominal_diameter = 48*mm,
    thickness = 38*mm,
    inscribed_diameter = 75*mm,
    circumscribed_diameter = 82.6*mm,
    bearing_diameter = 69.4*mm)

M48_T = dict(
    nominal_diameter = 48*mm,
    thread_pitch = 5*mm,
    shank_diameter_ratio = 44.752/48,
    shank_length = 150*mm,
    yield_stress = 900*MPa,
    ultimate_tensile_stress = 1000*MPa,
    stud = True)

M48_T_hex_nut = dict(
    nominal_diameter = 48*mm,
    thickness = 64*mm,
    inscribed_diameter = 75*mm,
    circumscribed_diameter = 82.6*mm,
    bearing_diameter = 92*mm)


D7500_L = dict(
    flange_type = "L",
    D = 7500*mm,
    n_bolts = 120,
    a = 232.5*mm,
    b = 166.5*mm,
    s = 72*mm,
    t = 200*mm,
    Zg = -14795*kN / 120,
    Fv = 2800*kN,
    Do = 86*mm,
    s_ratio = 100/72,
    bolt = M80,
    nut = M80_hex_nut,
    washer = None)

D4600_L = dict(
    flange_type = "L",
    D = 4600*mm,
    n_bolts = 144,
    a = 160*mm,
    b = 80*mm,
    s = 30*mm,
    t = 125*mm,
    Zg = -30.139*kN,
    Fv = 910*kN,
    Do = 52*mm,
    s_ratio = 30/30,
    bolt = M48_L,
    nut = M48_L_hex_nut,
    washer = dict(outer_diameter=92*mm, inner_diameter=49.4*mm, thickness=8*mm))

D7500_T = dict(
    flange_type = "T",
    D = 7500*mm,
    n_bolts = 200,
    a = 62.5*mm,
    b = 111.0*mm,
    s = 90*mm,
    t = 120*mm,
    Zg = -81.4*kN * 1.3,    # includes the macro-geometric factor 1.3
    Fv = 928*kN,
    Do = 52*mm,
    s_ratio = 1.0,
    bolt = M48_T,
    nut = M48_T_hex_nut,
    washer = None)

# Fatigue analysis parameters.
# L-flanges: bending factor according to IEC 61400-6 AMD1 (evaluated from the bolt diameter).
# T-flanges: as in Case-D7500_T.py, with macro-geometric factor 1.3.
L_FATIGUE = dict(bending_factor="IEC", mean_factor=1.0, range_factor=1.0)
T_FATIGUE = dict(bending_factor=0.601, mean_factor=1.3, range_factor=1.19*1.3)


STIFF_TAG = {"simplified": "simpl", "interpolated": "interp"}


def L_case (num, segment, shape_factor, tilt_deg, stiff, title):
    suffix = STIFF_TAG[stiff].capitalize()
    fname = {
        1: "Case_1_D7500_L_Tilt-0p00deg_ShapeFactor-1p00_Fv2800kN_ShellStiff-{}.xlsx",
        2: "Case_2_D7500_L_Tilt-0p25deg_ShapeFactor-1p00_Fv2800kN_ShellStiff-{}.xlsx",
        3: "Case_3_D7500_L_Tilt-0p00deg_ShapeFactor-1p20_Fv2800kN_ShellStiff-{}.xlsx",
        4: "Case_4_D4600_L_Tilt-0p00deg_ShapeFactor-1p00_Fv910kN_ShellStiff-{}.xlsx",
    }[num].format(suffix)
    return dict(
        id = f"L{num}-{STIFF_TAG[stiff]}",
        title = f"Case {num} - {title}, {stiff} shell stiffness",
        workbook = os.path.join(L_DIR, fname),
        sheet = "SGRE",
        first_col = 12,     # column L
        quantities = L_QUANTITIES,
        segment = segment,
        gap = dict(flatness_tolerance=0.0014, fractile=0.95, shape_factor=shape_factor),
        tilt_angle_deg = tilt_deg,
        shell_stiffness = stiff,
        fatigue = L_FATIGUE)


def T_case (num, shape_factor, tilt_deg, stiff, title):
    fname = {
        5: "Case_5-D7500_Tilt-0deg_ShapeFactor-1.0{}.xlsx",
        6: "Case_6-D7500_Tilt-0deg_ShapeFactor-1.5{}.xlsx",
        7: "Case_7-D7500_Tilt-1deg_ShapeFactor-1.5{}.xlsx",
        8: "Case_8-D7500_Tilt-1deg_ShapeFactor-1.0{}.xlsx",
    }[num].format("-interpolated" if stiff == "interpolated" else "")
    return dict(
        id = f"T{num}-{STIFF_TAG[stiff]}",
        title = f"Case {num} - {title}, {stiff} shell stiffness",
        workbook = os.path.join(T_DIR, fname),
        sheet = "AMD1",
        first_col = 14,     # column N
        quantities = T_QUANTITIES,
        segment = D7500_T,
        gap = dict(flatness_tolerance=0.0014, fractile=0.95, shape_factor=shape_factor),
        tilt_angle_deg = tilt_deg,
        shell_stiffness = stiff,
        fatigue = T_FATIGUE)


CASES = []
for stiff in ("simplified", "interpolated"):
    CASES += [
        L_case(1, D7500_L, 1.0, 0.00, stiff, "D7500 L-flange, sinusoidal gap, no tilt"),
        L_case(2, D7500_L, 1.0, 0.25, stiff, "D7500 L-flange, sinusoidal gap, 0.25 deg tilt"),
        L_case(3, D7500_L, 1.2, 0.00, stiff, "D7500 L-flange, gap shape factor 1.2, no tilt"),
        L_case(4, D4600_L, 1.0, 0.00, stiff, "D4600 L-flange, sinusoidal gap, no tilt"),
    ]
for stiff in ("simplified", "interpolated"):
    CASES += [
        T_case(5, 1.0, 0.0, stiff, "D7500 T-flange, sinusoidal gap, no tilt"),
        T_case(6, 1.5, 0.0, stiff, "D7500 T-flange, gap shape factor 1.5, no tilt"),
        T_case(7, 1.5, 1.0, stiff, "D7500 T-flange, gap shape factor 1.5, 1 deg tilt"),
        T_case(8, 1.0, 1.0, stiff, "D7500 T-flange, sinusoidal gap, 1 deg tilt"),
    ]



# ------------------------------------------------------------------------------
#  Markov matrices
# ------------------------------------------------------------------------------

def read_workbook_markov_matrix (wb):
    '''Reads the 'SGRE!markov_matrix' range of an L-flange workbook.

    The range has three columns: moment range [kNm], moment mean [kNm] and
    number of cycles. The load duration is 25 years.
    '''
    ref = wb["SGRE"].defined_names["markov_matrix"].attr_text
    cells = wb["SGRE"][ref.split("!")[1].replace("$", "")]
    rows = [[c.value for c in row] for row in cells]
    rows = [r for r in rows if all(isinstance(v, (int, float)) for v in r)]
    return dict(
        range = [r[0]*1e3 for r in rows],
        mean = [r[1]*1e3 for r in rows],
        cycles = [r[2] for r in rows],
        duration = 25)


def read_mkv_markov_matrix (path, unit_factor=1e3, duration=1):
    '''Reads a Markov matrix from a SGRE .mkv file.

    Port of `markov_matrix_from_SGRE_format` from Case-D7500_T.py. The values
    in the file are in kNm, therefore they get converted to Nm.
    '''
    with open(path) as mkv_file:
        lines = mkv_file.readlines()

    mkv_range, mkv_mean, mkv_cycles = [], [], []
    separators = 0
    mean_values = None
    for line in lines:
        if line.strip() == "---------------------------":
            separators += 1
            continue
        if separators < 2:
            continue
        values = line.split()
        if not values:
            continue
        if mean_values is None:
            mean_values = values
            continue
        for i in range(1, len(values)):
            if float(values[i]) == 0.0: continue
            mkv_cycles.append(float(values[i]))
            mkv_mean.append(float(mean_values[i]) * unit_factor)
            mkv_range.append(float(values[0]) * unit_factor)

    return dict(range=mkv_range, mean=mkv_mean, cycles=mkv_cycles, duration=duration)


def save_markov_matrix (mkvm, name_hint, registry):
    '''Saves the Markov matrix once per distinct content and returns its relative path.'''
    digest = hashlib.sha1(json.dumps(mkvm, sort_keys=True).encode()).hexdigest()
    if digest not in registry:
        rel_path = f"markov/{name_hint}.json"
        n = 2
        while rel_path in registry.values():
            rel_path = f"markov/{name_hint}-{n}.json"
            n += 1
        os.makedirs(os.path.join(OUT_DIR, "markov"), exist_ok=True)
        with open(os.path.join(OUT_DIR, rel_path), "w") as f:
            json.dump(mkvm, f)
        registry[digest] = rel_path
    return registry[digest]



# ------------------------------------------------------------------------------
#  Main
# ------------------------------------------------------------------------------

def read_references (ws, quantities, first_col):
    references = {}
    for key, (row, symbol, unit) in sorted(quantities.items(), key=lambda item: item[1][0]):
        values = [ws.cell(row, first_col + i).value for i in range(len(GAP_ANGLES))]
        if any(not isinstance(v, (int, float)) for v in values):
            raise ValueError(f"Non-numeric reference value for '{key}' at row {row}: {values}")
        references[key] = dict(symbol=symbol, unit=unit, sgre_row=row, values=values)
    return references


def main ():
    os.makedirs(OUT_DIR, exist_ok=True)
    markov_registry = {}
    t_markov = read_mkv_markov_matrix(os.path.join(T_DIR, "tflange-example-markov.mkv"))

    for case in CASES:
        print(f"Extracting {case['id']} from {os.path.relpath(case['workbook'], HERE)}")
        wb = openpyxl.load_workbook(case["workbook"], data_only=True)
        ws = wb[case["sheet"]]

        # Check that the gap angles in the workbook are the expected ones
        angles = [ws.cell(48, case["first_col"] + i).value for i in range(len(GAP_ANGLES))]
        assert angles == GAP_ANGLES, f"Unexpected gap angles {angles}"

        if case["segment"]["flange_type"] == "L":
            mkvm = read_workbook_markov_matrix(wb)
            hint = f"L-D{round(case['segment']['D']/mm)}"
        else:
            mkvm = t_markov
            hint = "T-D7500"

        data = dict(
            id = case["id"],
            title = case["title"],
            source = os.path.relpath(case["workbook"], HERE).replace(os.sep, "/"),
            segment = case["segment"],
            gap = case["gap"],
            tilt_angle_deg = case["tilt_angle_deg"],
            shell_stiffness = case["shell_stiffness"],
            fatigue = dict(markov_matrix=save_markov_matrix(mkvm, hint, markov_registry), **case["fatigue"]),
            gap_angles_deg = GAP_ANGLES,
            references = read_references(ws, case["quantities"], case["first_col"]))

        # Keep the tolerances already set in the case file (they are not in the workbooks)
        out_path = os.path.join(OUT_DIR, f"{case['id']}.json")
        if os.path.exists(out_path):
            with open(out_path, encoding="utf-8") as f:
                old_data = json.load(f)
            if "tolerances" in old_data:
                data["tolerances"] = old_data["tolerances"]

        with open(out_path, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2, ensure_ascii=False)


if __name__ == "__main__":
    main()
